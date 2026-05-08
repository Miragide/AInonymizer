"""
Extraction de texte brut depuis PDF, DOCX et Markdown.
- PDF  : pymupdf/fitz (MuPDF, sans subprocess, sans OCR)
- DOCX : python-docx avec préservation des niveaux de titres
- MD   : lecture directe UTF-8
"""
from __future__ import annotations

import re
from pathlib import Path


# ═══════════════════════════════════════════════════════════════════════
#  Point d'entrée public
# ═══════════════════════════════════════════════════════════════════════

def extract(path: Path) -> str:
    """Retourne le texte brut du fichier. Lève ValueError si format inconnu."""
    ext = path.suffix.lower()
    if ext == ".pdf":
        return _extract_pdf(path)
    if ext == ".docx":
        return _extract_docx(path)
    if ext in (".md", ".markdown"):
        return path.read_text(encoding="utf-8", errors="replace")
    raise ValueError(f"Format non supporté : {ext}")


# ═══════════════════════════════════════════════════════════════════════
#  PDF — pymupdf (MuPDF)
# ═══════════════════════════════════════════════════════════════════════

def _extract_pdf(path: Path) -> str:
    import fitz  # pymupdf
    doc = fitz.open(str(path))
    parts = [page.get_text() for page in doc]
    doc.close()
    return _clean_pdf("\n\n".join(parts))


_FF = re.compile(r"\f")                          # saut de page
_MULTI_BLANK = re.compile(r"\n{3,}")             # trop de lignes vides
_TRAILING_WS = re.compile(r"[ \t]+$", re.MULTILINE)
_MULTI_SPACE = re.compile(r"  +")               # espaces multiples dans une ligne


def _clean_pdf(text: str) -> str:
    text = _FF.sub("\n\n", text)
    text = _TRAILING_WS.sub("", text)
    text = _MULTI_SPACE.sub(" ", text)
    text = _MULTI_BLANK.sub("\n\n", text)
    return text.strip()


# ═══════════════════════════════════════════════════════════════════════
#  DOCX — python-docx
# ═══════════════════════════════════════════════════════════════════════

def _extract_docx(path: Path) -> str:
    from docx import Document

    doc = Document(str(path))
    parts: list[str] = []

    for para in doc.paragraphs:
        t = para.text.strip()
        if not t:
            continue
        style = para.style.name if para.style else ""
        if style.startswith("Heading") and style[-1].isdigit():
            level = int(style[-1])
            parts.append("#" * level + " " + t)
        else:
            parts.append(t)

    # Tables
    for table in doc.tables:
        rows: list[list[str]] = []
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            rows.append(cells)
        if not rows:
            continue
        # Ligne d'en-tête
        header = " | ".join(rows[0])
        sep = " | ".join("---" for _ in rows[0])
        parts.append(header)
        parts.append(sep)
        for row in rows[1:]:
            parts.append(" | ".join(row))

    return "\n\n".join(parts)
